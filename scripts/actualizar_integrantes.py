#!/usr/bin/env python3
"""Actualiza la identificación grupal de la PC2 conservando las fuentes LEC06/07 intactas."""
from pathlib import Path
from docx import Document
from docx.shared import Pt
import re
import shutil
import subprocess
import tempfile

ROOT = Path(__file__).resolve().parents[1]
MEMBERS = [
    ("Santa Cruz Basilio, Jesus Diego", "20235005G"),
    ("López Arteaga, César Omar", "20236743A"),
    ("Pumayauli Evangelista, Bryan José", "20212209E"),
]
TEAM_TEXT = " | ".join(f"{name} — {code}" for name, code in MEMBERS)
TEAM_BANNER = f"> **Integrantes, Grupo 10 — CC451 PC2:** {TEAM_TEXT}.\n\n"
REPORT = ROOT / "docs/01-informe-PC2.docx"
PERSONAL = ROOT / "mi-aporte-cesar-lopez-arteaga.docx"

def set_paragraph(paragraph, text):
    paragraph.text = text
    for run in paragraph.runs:
        run.font.size = Pt(11)

def update_report():
    doc = Document(REPORT)
    changed = 0
    for p in doc.paragraphs:
        if p.text.startswith("Estudiante: López Arteaga"):
            set_paragraph(p, "Integrantes del Grupo 10:")
            changed += 1
        elif p.text.startswith("Código: 20215508C"):
            set_paragraph(p, "\n".join(f"• {name} — {code}" for name, code in MEMBERS))
            changed += 1
    if changed not in (0, 2):
        raise RuntimeError(f"Portada del informe: actualización parcial ({changed})")
    for section in doc.sections:
        for paragraph in section.footer.paragraphs:
            if "César Omar López Arteaga" in paragraph.text:
                paragraph.text = paragraph.text.replace("César Omar López Arteaga", "Grupo 10: Santa Cruz · López · Pumayauli")
    doc.core_properties.author = "Grupo 10: Santa Cruz Basilio; López Arteaga; Pumayauli Evangelista"
    doc.core_properties.title = "PC2 CC451 — TriageVoz Rural — Grupo 10"
    doc.save(REPORT)

def update_personal():
    doc = Document(PERSONAL)
    ps = doc.paragraphs
    if not (ps[1].text.startswith("Estudiante:") or ps[1].text.startswith("Grupo 10")):
        raise RuntimeError("No se encontró identificación de los integrantes en el aporte")
    set_paragraph(ps[1], "Grupo 10 — integrantes:")
    set_paragraph(ps[2], "\n".join(f"• {name} — {code}" for name, code in MEMBERS))
    if not any(p.text.startswith("Aporte personal registrado:") for p in ps):
        ps[3].insert_paragraph_before("Aporte personal registrado: López Arteaga, César Omar — 20236743A.")
    for p in doc.paragraphs:
        p.text = p.text.replace("No se atribuyen contribuciones de integrantes del grupo no identificados.",
            "Los tres integrantes están identificados; el aporte individual descrito debe verificarse por el estudiante antes de la sustentación.")
        p.text = p.text.replace("20215508C", "20236743A")
    doc.core_properties.author = "Grupo 10 (aporte individual de López Arteaga, César Omar)"
    doc.save(PERSONAL)

def update_markdown():
    docs = [ROOT / "README.md", ROOT / "mi-aporte-cesar-lopez-arteaga.md", *sorted((ROOT / "docs").glob("*.md"))]
    for path in docs:
        if not path.exists():
            continue
        t = path.read_text(encoding="utf-8")
        t = t.replace("**Autor de la entrega:** César Omar López Arteaga — código 20215508C.", "**Integrantes del proyecto:** " + TEAM_TEXT + ".")
        t = t.replace("20215508C", "20236743A")
        if "Integrantes, Grupo 10 — CC451 PC2:" not in t:
            if t.startswith("# "):
                first, sep, rest = t.partition("\n")
                t = first + "\n\n" + TEAM_BANNER + rest.lstrip("\n")
            else:
                t = TEAM_BANNER + t
        path.write_text(t, encoding="utf-8")
    (ROOT / "docs/12-integrantes-grupo-10.md").write_text(
        "# Integrantes oficiales del Grupo 10 — PC2 CC451\n\n" +
        "\n".join(f"- **{name}** — {code}" for name, code in MEMBERS) +
        "\n\n**Proyecto:** TriageVoz Rural — Pre-triaje multimodal (PWA).\n\n" +
        "**Nota documental:** LEC06 y LEC07 se conservan intactos porque son fuentes previas, no documentos de autoría conjunta de PC2. "
        "El archivo mi-aporte-cesar-lopez-arteaga identifica al equipo, pero su descripción de labores sigue siendo personal, sin asignar trabajo a los otros integrantes.\n", encoding="utf-8"
    )

def convert(docx_path, out_dir, extension):
    lo = shutil.which("libreoffice") or shutil.which("soffice")
    if not lo:
        raise RuntimeError("Se requiere LibreOffice para generar los archivos DOC y PDF")
    with tempfile.TemporaryDirectory() as temp:
        cmd = [lo, "-env:UserInstallation=file:///tmp/pc2_group_lo", "--headless",
               "--convert-to", extension, "--outdir", temp, str(docx_path)]
        res = subprocess.run(cmd, capture_output=True, text=True, timeout=180)
        generated = Path(temp) / (docx_path.stem + "." + extension)
        if res.returncode or not generated.exists():
            raise RuntimeError("Falló conversión " + extension + ": " + res.stdout + res.stderr)
        target = out_dir / generated.name
        shutil.copy2(generated, target)
        return target

if __name__ == "__main__":
    update_report()
    update_personal()
    update_markdown()
    pdf = convert(REPORT, ROOT / "docs", "pdf")
    doc = convert(PERSONAL, ROOT, "doc")
    print("Archivos actualizados:", REPORT, pdf, PERSONAL, doc)
    print("Equipo:", TEAM_TEXT)
