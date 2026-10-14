#!/usr/bin/env python3
"""Convierte el anexo de revisión (Markdown) en DOCX reproducible. La conversión a PDF se realiza en CI."""
from pathlib import Path
from docx import Document
from docx.shared import Cm,Pt,RGBColor
import re
root=Path(__file__).resolve().parents[1]
source=root/'docs/14-anexo-version-actual.md'
target=root/'docs/14-anexo-version-actual.docx'
doc=Document()
section=doc.sections[0]
section.top_margin=Cm(2.3);section.bottom_margin=Cm(2.2)
section.left_margin=Cm(2.6);section.right_margin=Cm(2.4)
normal=doc.styles['Normal']
normal.font.name='Aptos';normal.font.size=Pt(10.5)
normal.paragraph_format.space_after=Pt(7)
for heading in ['Title','Heading 1','Heading 2']:
    doc.styles[heading].font.name='Aptos'
    doc.styles[heading].font.color.rgb=RGBColor(15,78,101)
for raw in source.read_text(encoding='utf-8').splitlines():
    line=raw.strip()
    if not line:continue
    if line.startswith('# '):doc.add_heading(line[2:],0)
    elif line.startswith('## '):doc.add_heading(line[3:],1)
    elif line.startswith('### '):doc.add_heading(line[4:],2)
    elif line.startswith('>'):doc.add_paragraph(line.lstrip('> ').replace('**',''))
    elif line.startswith('- '):doc.add_paragraph(line[2:].replace('**','').replace('`',''),'List Bullet')
    elif re.match(r'^\d+\. ',line):doc.add_paragraph(re.sub(r'^\d+\. ','',line).replace('**','').replace('`',''),'List Number')
    else:doc.add_paragraph(line.replace('**','').replace('`',''))
footer=section.footer.paragraphs[0]
footer.text='UNI · CC451 · Grupo 10 · Anexo de integración PC2 · 2026-II'
footer.style='Caption'
doc.core_properties.title='Anexo de integración PC2 - TriageVoz Rural'
doc.core_properties.author='Grupo 10 - Santa Cruz Basilio, López Arteaga, Pumayauli Evangelista'
target.parent.mkdir(parents=True,exist_ok=True)
doc.save(target)
# Incorporar una sola vez el contenido técnico al informe grupal, sin sustituir sus capítulos existentes.
report_path=root/'docs/01-informe-PC2.docx'
if report_path.exists():
    report=Document(report_path)
    annex_title='ANEXO DE ACTUALIZACIÓN: INTEGRACIÓN DEL GRUPO 10'
    if not any(annex_title in p.text for p in report.paragraphs):
        report.add_page_break()
        report.add_heading(annex_title,level=1)
        report.add_paragraph('Este anexo registra cambios posteriores al informe original y no modifica las fuentes LEC06/LEC07.')
        for paragraph in doc.paragraphs:
            value=paragraph.text.strip()
            if not value or value.startswith('Anexo técnico de integración') or value.startswith('Grupo 10 UNI:'):continue
            if paragraph.style.name=='Heading 1':
                report.add_heading(value,level=2)
            elif paragraph.style.name=='Heading 2':
                report.add_heading(value,level=3)
            else:
                report.add_paragraph(value)
        report.core_properties.author='Grupo 10 - Santa Cruz Basilio, López Arteaga, Pumayauli Evangelista'
        report.save(report_path)
        print('Informe DOCX grupal ampliado:',report_path)
    else:
        print('Anexo ya presente; no se duplica en el informe.')
print(target)
