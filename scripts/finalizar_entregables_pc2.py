#!/usr/bin/env python3
"""Normaliza códigos oficiales y regenera entregables PC2 desde fuentes verificables.
Las fuentes académicas originales LEC06/LEC07/PC02 se conservan intactas.
"""
from pathlib import Path
import subprocess
import re
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from zipfile import ZipFile
from PIL import Image

ROOT=Path(__file__).resolve().parents[1]
OLD='202'+'36743A'
NEW='20215508C'
NAMES=[
 ('Santa Cruz Basilio, Jesus Diego','20235005G'),
 ('López Arteaga, César Omar',NEW),
 ('Pumayauli Evangelista, Bryan José','20212209E'),
]
TEXT_EXTENSIONS={'.md','.csv','.py','.cjs','.js','.mjs','.html','.css','.yml','.yaml','.json','.txt','.dot','.svg','.webmanifest'}
EXCLUDED={'.git','node_modules','.venv','__pycache__'}

def run(cmd):
    print('+',' '.join(map(str,cmd)),flush=True)
    subprocess.run(list(map(str,cmd)),cwd=ROOT,check=True,timeout=240)

def update_sources():
    count=0
    for path in ROOT.rglob('*'):
        if not path.is_file() or path.suffix.lower() not in TEXT_EXTENSIONS:continue
        if any(x in path.parts for x in EXCLUDED):continue
        content=path.read_text(encoding='utf-8')
        new=content.replace(OLD,NEW)
        if path.name=='generar_presentacion.cjs':
            new=new.replace('Capturas actualizadas y video de sustentación de cada integrante.',
                'Capturas verificables disponibles; video personal de sustentación pendiente.')
        if path.name=='10-checklist-PC2.md':
            new=new.replace('assets/ (histórico), video no adjunto | **Pendiente**',
                'assets/evidencias-integracion (5 capturas E2E), video no adjunto | Capturas disponibles; video pendiente')
        if path.name=='00-ENTREGA-PC2.md':
            new=new.replace('El registro de screenshots de la interfaz final y el video siguen pendientes de captura y grabación por los integrantes; no se declaran como evidencia disponible antes de realizarlos.',
                'Las cinco capturas E2E del funcionamiento actual están publicadas y verificadas en GitHub Actions; el video de sustentación individual sigue pendiente.')
        if new!=content:
            path.write_text(new,encoding='utf-8')
            count+=1
    print('Fuentes de texto corregidas:',count)

def styled_runs(paragraph,line):
    line=line.replace('`','')
    parts=re.split(r'(\*\*.*?\*\*)',line)
    for part in parts:
        if not part:continue
        strong=part.startswith('**') and part.endswith('**')
        r=paragraph.add_run(part[2:-2] if strong else part)
        if strong:r.bold=True

def new_personal():
    source=(ROOT/'mi-aporte-cesar-lopez-arteaga.md').read_text(encoding='utf-8')
    doc=Document()
    sec=doc.sections[0]
    sec.top_margin=Cm(1.70);sec.bottom_margin=Cm(1.63)
    sec.left_margin=Cm(2.17);sec.right_margin=Cm(2.12)
    styles=doc.styles
    normal=styles['Normal'];normal.font.name='Aptos';normal.font.size=Pt(9.4)
    normal.paragraph_format.space_after=Pt(4.5)
    normal.paragraph_format.line_spacing=1.07
    for st in ('Title','Heading 1','Heading 2'):
        styles[st].font.name='Aptos'
        styles[st].font.color.rgb=RGBColor(18,71,112)
    styles['Title'].font.size=Pt(17)
    styles['Heading 1'].font.size=Pt(12)
    figure_definitions=[
      ('01-caso-enviado.png','Figura 1. Caso de pre-triaje registrado y enviado al servidor local.'),
      ('02-bandeja-casos.png','Figura 2. Bandeja del médico remoto con el caso disponible para consulta.'),
      ('03-confirmacion-historial.png','Figura 3. Acuse del médico simulado y respuesta conservada en el historial.'),
      ('04-respuesta-paciente.png','Figura 4. Confirmación recibida y mensaje visible desde la vista paciente.'),
    ]
    figure_section=False
    for raw in source.splitlines():
        line=raw.strip()
        if not line:continue
        if line.startswith('# '):
            doc.add_heading(line[2:],level=0)
        elif line.startswith('## '):
            if line.startswith('## 4.'):
                figure_section=True
                doc.add_page_break()
            doc.add_heading(line[3:],level=1)
        elif not figure_section:
            if line.startswith('- '):
                p=doc.add_paragraph(style='List Bullet')
                styled_runs(p,line[2:])
            else:
                p=doc.add_paragraph()
                styled_runs(p,line)
        elif figure_section and not line.startswith('- '):
            p=doc.add_paragraph()
            styled_runs(p,line)
    for i,(filename,caption) in enumerate(figure_definitions):
        image=ROOT/'assets/evidencias-integracion'/filename
        if not image.is_file():raise FileNotFoundError(image)
        if i>0:doc.add_page_break()
        p=doc.add_paragraph()
        p.alignment=WD_ALIGN_PARAGRAPH.CENTER
        p.add_run().add_picture(str(image),width=Cm(15.95))
        p=doc.add_paragraph()
        p.alignment=WD_ALIGN_PARAGRAPH.CENTER
        r=p.add_run(caption);r.bold=True;r.font.size=Pt(9)
        p=doc.add_paragraph('Fuente: ejecución automatizada del prototipo TriageVoz Rural en Chromium, datos ficticios. Sesión distinta de la prueba manual SALA-555CEF76E974.')
        p.style='Caption'
        p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    footer=sec.footer.paragraphs[0]
    footer.alignment=WD_ALIGN_PARAGRAPH.CENTER
    footer.text='López Arteaga, César Omar — '+NEW+'  |  PC2 CC451 · Grupo 10'
    footer.runs[0].font.size=Pt(8)
    doc.core_properties.author='López Arteaga, César Omar'
    doc.core_properties.title='Aporte personal PC2 CC451 y evidencia visual - César Omar López Arteaga'
    output=ROOT/'mi-aporte-cesar-lopez-arteaga.docx'
    doc.save(output)
    assert len(Document(output).inline_shapes)==4
    print('Aporte individual primera persona y 4 capturas:',output)

def all_paragraphs(doc):
    def iter_table(t):
        for row in t.rows:
            for cell in row.cells:
                for p in cell.paragraphs:yield p
                for sub in cell.tables:yield from iter_table(sub)
    for p in doc.paragraphs:yield p
    for t in doc.tables:yield from iter_table(t)
    for sec in doc.sections:
        for area in [sec.header,sec.footer,sec.first_page_header,sec.first_page_footer]:
            for p in area.paragraphs:yield p
            for t in area.tables:yield from iter_table(t)

def fix_other_docx():
    for path in ROOT.rglob('*.docx'):
        if path.name=='mi-aporte-cesar-lopez-arteaga.docx':continue
        doc=Document(path)
        changes=0
        for p in all_paragraphs(doc):
            if OLD not in p.text:continue
            before=p.text
            for run in p.runs:
                run.text=run.text.replace(OLD,NEW)
            if OLD in p.text:
                p.text=before.replace(OLD,NEW)
            changes+=1
        if changes:
            doc.save(path)
            print('Código corregido en Word:',path.relative_to(ROOT),changes)

def convert(src,dest_dir,format):
    dest_dir.mkdir(parents=True,exist_ok=True)
    out=dest_dir/(src.stem+'.'+format)
    if out.exists():out.unlink()
    run(['libreoffice','-env:UserInstallation=file:///tmp/triagevoz_final_2026',
         '--headless','--convert-to',format,'--outdir',dest_dir,src])
    if not out.is_file() or out.stat().st_size<3000:
        raise RuntimeError('Error al convertir '+str(src)+' a '+format)
    print('Exportado:',out.relative_to(ROOT),out.stat().st_size)

def export():
    for stem in ('mi-aporte-cesar-lopez-arteaga','mi-aporte-jesus-santa-cruz','mi-aporte-bryan-pumayauli'):
        src=ROOT/(stem+'.docx')
        if src.exists():convert(src,ROOT,'doc')
    for stem in ('01-informe-PC2','14-anexo-version-actual'):
        src=ROOT/'docs'/(stem+'.docx')
        if src.exists():convert(src,ROOT/'docs','pdf')
    convert(ROOT/'presentacion/PC2-TriageVoz-Rural-Grupo10.pptx',ROOT/'presentacion','pdf')

def verify():
    for path in ROOT.rglob('*.md'):
        if OLD in path.read_text(encoding='utf-8'):raise RuntimeError('Código anterior en '+str(path))
    for path in [ROOT/'mi-aporte-cesar-lopez-arteaga.docx',ROOT/'docs/01-informe-PC2.docx',
                 ROOT/'mi-aporte-jesus-santa-cruz.docx',ROOT/'mi-aporte-bryan-pumayauli.docx']:
        doc=Document(path)
        content=' '.join(p.text for p in all_paragraphs(doc))
        if OLD in content:raise RuntimeError('Código antiguo en '+str(path))
    cs=Document(ROOT/'mi-aporte-cesar-lopez-arteaga.docx')
    full=' '.join(p.text for p in all_paragraphs(cs))
    if NEW not in full or len(cs.inline_shapes)!=4 or 'Consolidé' not in full:
        raise RuntimeError('Aporte César incompleto')
    ppt=ROOT/'presentacion/PC2-TriageVoz-Rural-Grupo10.pptx'
    with ZipFile(ppt) as z:
        xml=' '.join(z.read(p).decode('utf-8') for p in z.namelist() if p.startswith('ppt/slides/') and p.endswith('.xml'))
        if OLD in xml or NEW not in xml:raise RuntimeError('Presentación contiene código incorrecto')
    print('VERIFICACIÓN: códigos oficiales y cuatro capturas incrustadas, correctos.')

if __name__=='__main__':
    update_sources()
    new_personal()
    fix_other_docx()
    run(['node','scripts/generar_presentacion.cjs'])
    export()
    verify()
